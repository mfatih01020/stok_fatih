import os

with open('app.py', 'rb') as f:
    content = f.read().decode('utf-8', errors='replace')

lines = content.splitlines()

# Let's inspect where api_compare starts and download_remaining ends
start_idx = None
end_idx = None
for i, line in enumerate(lines):
    if '@app.route(\'/api/compare\'' in line:
        start_idx = i
    if '@app.route(\'/api/audit_box\'' in line:
        end_idx = i

print(f'api_compare start: {start_idx}, audit_box start: {end_idx}')

new_compare_block = """@app.route('/api/compare', methods=['POST'])
def api_compare():
    try:
        if 'system_file' not in request.files or 'sales_file' not in request.files:
            return jsonify({"success": False, "error": "Lütfen hem Sistem Depo hem de Satış dosyasını yükleyin."})
            
        system_file = request.files['system_file']
        sales_file = request.files['sales_file']
        
        gtin_map = build_gtin_name_map()
        
        inventory = parse_system_file(system_file, gtin_map)
        sales_qrs = parse_sales_file(sales_file)
        
        if not inventory:
            return jsonify({"success": False, "error": "Sistem Depo dosyasından geçerli karekod okunamadı."})
        if not sales_qrs:
            return jsonify({"success": False, "error": "Satış dosyasından geçerli karekod okunamadı."})
            
        inventory_dict = {item["qr"]: item for item in inventory}
        matched_sales = []
        unmatched_sales = []
        sold_qrs_set = set()
        sales_by_product = {}
        
        # --- KOLI (BOX) STATUS ANALYSIS ---
        koli_stats = {}
        for item in inventory:
            meta = item.get("metadata", {})
            koli_no = meta.get("Koli Numarası") or meta.get("Paket Numarası") or meta.get("Palet Numarası") or meta.get("KOLINO")
            if koli_no:
                koli_no = str(koli_no).strip()
                if koli_no not in koli_stats:
                    koli_stats[koli_no] = {"total": 0, "sold": 0, "product_name": item["product_name"]}
                koli_stats[koli_no]["total"] += 1
        
        for sale_qr in sales_qrs:
            if sale_qr in inventory_dict:
                item = dict(inventory_dict[sale_qr])
                meta = item.get("metadata", {})
                koli_no = meta.get("Koli Numarası") or meta.get("Paket Numarası") or meta.get("Palet Numarası") or meta.get("KOLINO")
                if koli_no:
                    koli_no = str(koli_no).strip()
                    if koli_no in koli_stats:
                        koli_stats[koli_no]["sold"] += 1

                item["aciklama"] = "Bu ürün sistemde görünüyor, sistemden çık"
                matched_sales.append(item)
                sold_qrs_set.add(sale_qr)
                p_name = item["product_name"]
                sales_by_product[p_name] = sales_by_product.get(p_name, 0) + 1
            else:
                gtin = extract_gtin(sale_qr)
                p_name = gtin_map.get(gtin, "Sistem Dışı Ürün (GTIN: {})".format(gtin) if gtin else "Bilinmeyen Karekod formatı")
                unmatched_sales.append({
                    "qr": sale_qr,
                    "product_name": p_name,
                    "aciklama": "Depoda bulunamadı (sistem dışı satış)"
                })
                sales_by_product[p_name] = sales_by_product.get(p_name, 0) + 1
                
        # Assign detailed koli status to matched items
        for item in matched_sales:
            meta = item.get("metadata", {})
            koli_no = meta.get("Koli Numarası") or meta.get("Paket Numarası") or meta.get("Palet Numarası") or meta.get("KOLINO")
            if koli_no:
                koli_no = str(koli_no).strip()
                stats = koli_stats.get(koli_no)
                if stats:
                    sold_cnt = stats["sold"]
                    tot_cnt = stats["total"]
                    rem_cnt = max(0, tot_cnt - sold_cnt)
                    if sold_cnt >= tot_cnt:
                        item["Koli Durumu"] = f"Koli {koli_no}: Tamamı Satıldı ({sold_cnt}/{tot_cnt})"
                    else:
                        item["Koli Durumu"] = f"Koli {koli_no}: Kısmi Satıldı ({sold_cnt}/{tot_cnt} - Depoda {rem_cnt} Kaldı)"
            else:
                item["Koli Durumu"] = "Koli Bilgisi Yok"
                
        remaining_inventory = [item for item in inventory if item["qr"] not in sold_qrs_set]
        
        remaining_counts = {}
        for item in remaining_inventory:
            p_name = item["product_name"]
            remaining_counts[p_name] = remaining_counts.get(p_name, 0) + 1
            
        initial_counts = {}
        for item in inventory:
            p_name = item["product_name"]
            initial_counts[p_name] = initial_counts.get(p_name, 0) + 1

        global cached_results
        cached_results = {
            "matched_sales": matched_sales,
            "unmatched_sales": unmatched_sales,
            "remaining_inventory": remaining_inventory,
            "sales_by_product": sales_by_product,
            "initial_counts": initial_counts,
            "remaining_counts": remaining_counts,
            "koli_stats": koli_stats
        }

        return jsonify({
            "success": True,
            "total_initial": len(inventory),
            "total_sold": len(sales_qrs),
            "total_matched": len(matched_sales),
            "total_unmatched": len(unmatched_sales),
            "total_remaining": len(remaining_inventory),
            "matched_sales": matched_sales,
            "sales_by_product": sales_by_product,
            "remaining_counts": remaining_counts
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route('/api/download/sales', methods=['GET'])
@app.route('/api/download/full_report', methods=['GET'])
def download_sales():
    global cached_results
    if 'cached_results' not in globals() or not cached_results:
        return "No comparison run yet", 400
        
    # Sheet 1: Matches (Sistemden Düşülecek Satışlar)
    matched_rows = []
    for item in cached_results["matched_sales"]:
        meta = item.get("metadata", {})
        koli_no = meta.get("Koli Numarası") or meta.get("Paket Numarası") or meta.get("Palet Numarası") or meta.get("KOLINO") or ""
        gtin = extract_gtin(item["qr"]) or meta.get("Gtin Numarası") or meta.get("BARKOD") or ""
        seri_no = meta.get("Seri Numarası") or meta.get("SERINO") or ""
        parti_no = meta.get("Parti Numarası") or meta.get("SARJNO") or ""
        palet_no = meta.get("Palet Numarası") or meta.get("PALETNO") or ""
        uretim = meta.get("Üretim Tarihi") or meta.get("URETIMTARIHI") or ""
        skt = meta.get("Son Kullanma Tarihi") or meta.get("SKT") or ""

        row = {
            "Ürün Adı": item["product_name"],
            "Karekod": item["qr"],
            "Gtin / Barkod": gtin,
            "Koli Numarası": koli_no,
            "Seri Numarası": seri_no,
            "Parti Numarası": parti_no,
            "Palet Numarası": palet_no,
            "Üretim Tarihi": uretim,
            "Son Kullanma Tarihi": skt,
            "Koli / Depo Durumu": item.get("Koli Durumu", ""),
            "Açıklama / İşlem": "Bu ürün sistemde görünüyor, sistemden çık"
        }
        matched_rows.append(row)

    df_matched = pd.DataFrame(matched_rows)
    if not df_matched.empty:
        df_matched = df_matched.sort_values(by=["Ürün Adı"])
        
    # Sheet 2: Unmatched (Sistemde Olmayan Hatalı Satışlar)
    unmatched_rows = []
    for item in cached_results["unmatched_sales"]:
        unmatched_rows.append({
            "Ürün Adı": item["product_name"],
            "Karekod": item["qr"],
            "Açıklama / Durum": "Depoda Bulunamadı (Sistem Dışı Hatalı Satış!)"
        })
    df_unmatched = pd.DataFrame(unmatched_rows)
    if not df_unmatched.empty:
        df_unmatched = df_unmatched.sort_values(by=["Ürün Adı"])
        
    # Sheet 3: Summary (Satış Adetleri Özet)
    summary_rows = []
    for p_name, qty in cached_results["sales_by_product"].items():
        summary_rows.append({
            "Ürün Adı": p_name,
            "Satılan Miktar (Adet)": qty
        })
    df_summary = pd.DataFrame(summary_rows)
    if not df_summary.empty:
        df_summary = df_summary.sort_values(by=["Ürün Adı"])
        
    # Sheet 4: Koli Durum Özeti
    koli_rows = []
    for koli_no, stats in cached_results.get("koli_stats", {}).items():
        if stats["sold"] > 0:
            tot = stats["total"]
            sold = stats["sold"]
            rem = max(0, tot - sold)
            status = f"Tamamı Satıldı ({sold}/{tot})" if sold >= tot else f"Kısmi Satıldı ({sold}/{tot} - Depoda {rem} Kaldı)"
            koli_rows.append({
                "Koli Numarası": koli_no,
                "Ürün Adı": stats["product_name"],
                "Toplam Ürün": tot,
                "Satılan Ürün": sold,
                "Depoda Kalan": rem,
                "Durum": status
            })
    df_koli = pd.DataFrame(koli_rows)
    if not df_koli.empty:
        df_koli = df_koli.sort_values(by=["Durum", "Koli Numarası"])
        
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if not df_matched.empty:
            df_matched.to_excel(writer, index=False, sheet_name="Düşülecek Satışlar")
            
            # --- Excel Renklendirme ---
            try:
                from openpyxl.styles import PatternFill
                worksheet = writer.sheets["Düşülecek Satışlar"]
                fill_tamami = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Açık Yeşil
                fill_kismi = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Açık Sarı
                
                status_col_idx = None
                for col_num, col_name in enumerate(df_matched.columns, start=1):
                    if col_name == "Koli / Depo Durumu":
                        status_col_idx = col_num
                        break
                
                if status_col_idx:
                    for row_idx, status_val in enumerate(df_matched["Koli / Depo Durumu"], start=2):
                        if "Tamamı" in str(status_val):
                            for col_idx in range(1, len(df_matched.columns) + 1):
                                worksheet.cell(row=row_idx, column=col_idx).fill = fill_tamami
                        elif "Kısmi" in str(status_val):
                            for col_idx in range(1, len(df_matched.columns) + 1):
                                worksheet.cell(row=row_idx, column=col_idx).fill = fill_kismi
            except Exception as e:
                pass
                
        else:
            pd.DataFrame(columns=["Ürün Adı", "Karekod", "Koli Numarası", "Koli / Depo Durumu", "Açıklama / İşlem"]).to_excel(writer, index=False, sheet_name="Düşülecek Satışlar")
            
        if not df_unmatched.empty:
            df_unmatched.to_excel(writer, index=False, sheet_name="Depoda Olmayan Satışlar")
        if not df_koli.empty:
            df_koli.to_excel(writer, index=False, sheet_name="Koli Durum Özeti")
        if not df_summary.empty:
            df_summary.to_excel(writer, index=False, sheet_name="Satış Özet Tablosu")
            
    output.seek(0)
    
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="satis_raporu_detayli.xlsx"
    )

@app.route('/api/download/remaining', methods=['GET'])
def download_remaining():
    global cached_results
    if 'cached_results' not in globals() or not cached_results:
        return "No comparison run yet", 400
        
    rows = []
    for item in cached_results["remaining_inventory"]:
        meta = item.get("metadata", {})
        koli_no = meta.get("Koli Numarası") or meta.get("Paket Numarası") or meta.get("Palet Numarası") or meta.get("KOLINO") or ""
        gtin = extract_gtin(item["qr"]) or meta.get("Gtin Numarası") or meta.get("BARKOD") or ""
        seri_no = meta.get("Seri Numarası") or meta.get("SERINO") or ""
        parti_no = meta.get("Parti Numarası") or meta.get("SARJNO") or ""
        palet_no = meta.get("Palet Numarası") or meta.get("PALETNO") or ""
        uretim = meta.get("Üretim Tarihi") or meta.get("URETIMTARIHI") or ""
        skt = meta.get("Son Kullanma Tarihi") or meta.get("SKT") or ""

        row = {
            "Ürün Adı": item["product_name"],
            "Karekod": item["qr"],
            "Gtin / Barkod": gtin,
            "Koli Numarası": koli_no,
            "Seri Numarası": seri_no,
            "Parti Numarası": parti_no,
            "Palet Numarası": palet_no,
            "Üretim Tarihi": uretim,
            "Son Kullanma Tarihi": skt,
            "Durum": "Depoda Mevcut"
        }
        rows.append(row)
        
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(by=["Ürün Adı"])
        
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Güncel Kalan Envanter")
    output.seek(0)
    
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="kalan_sistem_envanteri_guncel.xlsx"
    )
"""

clean_lines = lines[:start_idx] + new_compare_block.splitlines() + lines[end_idx:]

with open('app.py', 'w', encoding='utf-8') as f:
    f.write('\n'.join(clean_lines))

print('app.py compare and report routes updated.')
